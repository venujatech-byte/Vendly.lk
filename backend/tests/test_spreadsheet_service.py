from openpyxl import load_workbook

from app.services.spreadsheet_service import (
    export_inventory_workbook,
    export_ledger_workbook,
    parse_inventory_workbook,
)


class FakeSnapshot:
    def __init__(self, document_id, data):
        self.id = document_id
        self._data = data

    def to_dict(self):
        return self._data


class FakeCollection:
    def __init__(self, records):
        self.records = records

    def document(self, document_id):
        return FakeBusinessDocument(self.records[document_id])

    def stream(self):
        return [FakeSnapshot(document_id, data) for document_id, data in self.records.items()]


class FakeBusinessDocument:
    def __init__(self, collections):
        self.collections = collections

    def collection(self, name):
        return FakeCollection(self.collections.get(name, {}))


class FakeDatabase:
    def __init__(self, businesses):
        self.businesses = businesses

    def collection(self, name):
        assert name == "businesses"
        return FakeCollection(self.businesses)


def test_inventory_export_round_trips_media_and_variant_images():
    database = FakeDatabase({
        "business-1": {
            "categories": {
                "category-1": {"name": "Shoes", "description": "Footwear", "status": "active", "sortOrder": 1},
            },
            "products": {
                "product-1": {
                    "name": "Running Shoes",
                    "categoryName": "Shoes",
                    "hasSizes": True,
                    "costPriceMinor": 120000,
                    "sellingPriceMinor": 189900,
                    "weightGrams": 450,
                    "media": [{"type": "image", "url": "https://images.example/shoe.jpg", "path": ""}],
                    "variantSummaries": [{
                        "id": "variant-1", "size": "36", "sku": "SHOE-36", "barcode": "1234567890123",
                        "stockOnHand": 5, "stockReserved": 1, "stockAvailable": 4,
                        "costPriceMinor": 120000, "sellingPriceMinor": 189900,
                        "weightGrams": 450, "imageUrl": "https://images.example/shoe-36.jpg",
                    }],
                },
            },
        },
    })

    exported = export_inventory_workbook(database, "business-1")
    parsed = parse_inventory_workbook(exported)

    assert parsed["products"][0]["payload"]["media"][0]["url"] == "https://images.example/shoe.jpg"
    assert parsed["products"][0]["payload"]["weightKg"] == 0.45
    assert parsed["products"][0]["payload"]["variants"][0]["imageUrl"] == "https://images.example/shoe-36.jpg"
    assert parsed["products"][0]["payload"]["variants"][0]["stock"] == 5


def test_ledger_export_keeps_credit_debit_and_balance_columns():
    exported = export_ledger_workbook([{ 
        "createdAt": "2026-08-29T10:00:00+00:00", "reference": "VD-000001",
        "label": "Online order", "customerName": "Kamal", "description": "Watch",
        "paymentMethod": "cod", "paymentStatus": "unpaid", "status": "confirmed",
        "direction": "credit", "amountMinor": 250000, "balanceMinor": 250000,
    }])
    sheet = load_workbook(exported, data_only=True)["Transaction Ledger"]

    assert sheet["B2"].value == "VD-000001"
    assert sheet["I2"].value == 2500
    assert sheet["J2"].value is None
    assert sheet["K2"].value == 2500
