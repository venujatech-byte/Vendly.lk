import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_security_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    pillars = [
        {
            "num": "01",
            "title": "File Upload\nImmunity",
            "bullets": [
                "Binary Magic Byte checking (JPEG, PNG, WebP, MP4)",
                "Disguised script & SVG payload prohibition",
                "Isolated cloud CDN storage (zero local execution)",
            ],
            "accent": RGBColor(6, 182, 212), # Cyan
        },
        {
            "num": "02",
            "title": "Search & XSS\nDefense",
            "bullets": [
                "HTML & script tag boundary stripping (<...>)",
                "Null byte (\\x00) & control char neutralization",
                "Bounded query length & ReDoS-immune linear matching",
            ],
            "accent": RGBColor(59, 130, 246), # Blue
        },
        {
            "num": "03",
            "title": "Instant Session\nRevocation",
            "bullets": [
                "Cryptographically verified Firebase JWT tokens",
                "Live token check (check_revoked=True)",
                "Zero zombie sessions on password reset or ban",
            ],
            "accent": RGBColor(99, 102, 241), # Indigo
        },
        {
            "num": "04",
            "title": "Enterprise\nRBAC Controls",
            "bullets": [
                "Granular role model (Admin, Orders, Staff, Viewer)",
                "Owner-protected hierarchy: peer admins cannot demote",
                "Strict multi-tenant store isolation",
            ],
            "accent": RGBColor(168, 85, 247), # Purple
        },
        {
            "num": "05",
            "title": "Perimeter\nDefense",
            "bullets": [
                "Global & tiered rate limiting (120 req/min)",
                "Clickjacking shield (X-Frame-Options: DENY)",
                "MIME-sniff protection (nosniff) & HSTS",
            ],
            "accent": RGBColor(16, 185, 129), # Emerald
        },
    ]

    speaker_notes = (
        "Security in Vendly.lk is built with a zero-trust, defense-in-depth architecture across 5 pillars:\n\n"
        "1. File Upload Immunity: We never rely on client MIME headers. True binary magic bytes (JPEG, PNG, WebP, MP4) "
        "are inspected, and disguised executable scripts are blocked. Media is served from isolated CDNs.\n"
        "2. Search & XSS Defense: All user queries strip script tags, HTML markers, and unprintable control characters. "
        "Strict length caps eliminate catastrophic ReDoS attacks.\n"
        "3. Live Session Revocation: Active JWTs are verified with check_revoked=True. Suspensions or password resets "
        "cut off access instantly, leaving no zombie tokens.\n"
        "4. Enterprise RBAC: Store owners maintain strict authority. Peer admins cannot escalate privileges, hijack ownership, "
        "or demote peer staff.\n"
        "5. Perimeter Defense: Global rate limiting blocks brute force, while enterprise HTTP headers prevent Clickjacking and MIME sniffing."
    )

    # -------------------------------------------------------------
    # SLIDE 1: Modern Dark Theme (Executive Tech)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)

    # Background
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
    bg1.line.fill.background()

    # Top accent bar
    top_bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(0.4), Inches(0.08))
    top_bar1.fill.solid()
    top_bar1.fill.fore_color.rgb = RGBColor(6, 182, 212)
    top_bar1.line.fill.background()

    # Title & Subtitle Box
    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(1.3))
    tf1 = title_box.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

    p_badge = tf1.paragraphs[0]
    p_badge.text = "VENDLY.LK PLATFORM SECURITY"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(6, 182, 212)

    p_title = tf1.add_paragraph()
    p_title.text = "Multi-Layered Threat Defense & Zero-Trust Architecture"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(248, 250, 252)
    p_title.space_after = Pt(4)

    p_sub = tf1.add_paragraph()
    p_sub.text = "End-to-end protection for merchant commerce, customer data, and system integrity."
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)

    # 5 Pillar Cards
    card_width = Inches(2.2)
    card_gap = Inches(0.18)
    start_x = Inches(0.8)
    card_y = Inches(2.15)
    card_height = Inches(4.7)

    for i, col in enumerate(pillars):
        x = start_x + i * (card_width + card_gap)
        
        # Card Background
        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(30, 41, 59) # Slate 800
        card.line.color.rgb = RGBColor(51, 65, 85) # Slate 700
        card.line.width = Pt(1.5)

        # Accent top strip inside card
        strip = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.12), card_y + Inches(0.12), card_width - Inches(0.24), Inches(0.06))
        strip.fill.solid()
        strip.fill.fore_color.rgb = col["accent"]
        strip.line.fill.background()

        # Text Frame inside card
        tb = slide1.shapes.add_textbox(x + Inches(0.15), card_y + Inches(0.28), card_width - Inches(0.3), card_height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = col["num"]
        p_num.font.size = Pt(13)
        p_num.font.bold = True
        p_num.font.color.rgb = col["accent"]
        p_num.space_after = Pt(2)

        p_h = tf.add_paragraph()
        p_h.text = col["title"]
        p_h.font.size = Pt(15)
        p_h.font.bold = True
        p_h.font.color.rgb = RGBColor(241, 245, 249)
        p_h.space_after = Pt(14)

        for bullet in col["bullets"]:
            p_b = tf.add_paragraph()
            p_b.text = f"•  {bullet}"
            p_b.font.size = Pt(10.5)
            p_b.font.color.rgb = RGBColor(203, 213, 225)
            p_b.space_after = Pt(10)

    # Attach speaker notes to Slide 1
    slide1.notes_slide.notes_text_frame.text = speaker_notes

    # -------------------------------------------------------------
    # SLIDE 2: Modern Light Theme (Clean Corporate)
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)

    # Background
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = RGBColor(248, 250, 252) # Slate 50
    bg2.line.fill.background()

    # Top accent bar
    top_bar2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(0.4), Inches(0.08))
    top_bar2.fill.solid()
    top_bar2.fill.fore_color.rgb = RGBColor(2, 132, 199)
    top_bar2.line.fill.background()

    # Title & Subtitle Box
    title_box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(1.3))
    tf2 = title_box2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0

    p2_badge = tf2.paragraphs[0]
    p2_badge.text = "VENDLY.LK PLATFORM SECURITY"
    p2_badge.font.size = Pt(11)
    p2_badge.font.bold = True
    p2_badge.font.color.rgb = RGBColor(2, 132, 199)

    p2_title = tf2.add_paragraph()
    p2_title.text = "Multi-Layered Threat Defense & Zero-Trust Architecture"
    p2_title.font.size = Pt(24)
    p2_title.font.bold = True
    p2_title.font.color.rgb = RGBColor(15, 23, 42)
    p2_title.space_after = Pt(4)

    p2_sub = tf2.add_paragraph()
    p2_sub.text = "End-to-end protection for merchant commerce, customer data, and system integrity."
    p2_sub.font.size = Pt(13)
    p2_sub.font.color.rgb = RGBColor(100, 116, 139)

    for i, col in enumerate(pillars):
        x = start_x + i * (card_width + card_gap)
        
        # Card Background
        card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_width, card_height)
        card2.fill.solid()
        card2.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card2.line.color.rgb = RGBColor(226, 232, 240)
        card2.line.width = Pt(1.5)

        # Accent top strip inside card
        strip2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.12), card_y + Inches(0.12), card_width - Inches(0.24), Inches(0.06))
        strip2.fill.solid()
        strip2.fill.fore_color.rgb = col["accent"]
        strip2.line.fill.background()

        # Text Frame inside card
        tb2 = slide2.shapes.add_textbox(x + Inches(0.15), card_y + Inches(0.28), card_width - Inches(0.3), card_height - Inches(0.4))
        tf_c2 = tb2.text_frame
        tf_c2.word_wrap = True
        tf_c2.margin_left = tf_c2.margin_right = tf_c2.margin_top = tf_c2.margin_bottom = 0

        p2_num = tf_c2.paragraphs[0]
        p2_num.text = col["num"]
        p2_num.font.size = Pt(13)
        p2_num.font.bold = True
        p2_num.font.color.rgb = col["accent"]
        p2_num.space_after = Pt(2)

        p2_h = tf_c2.add_paragraph()
        p2_h.text = col["title"]
        p2_h.font.size = Pt(15)
        p2_h.font.bold = True
        p2_h.font.color.rgb = RGBColor(15, 23, 42)
        p2_h.space_after = Pt(14)

        for bullet in col["bullets"]:
            p2_b = tf_c2.add_paragraph()
            p2_b.text = f"•  {bullet}"
            p2_b.font.size = Pt(10.5)
            p2_b.font.color.rgb = RGBColor(71, 85, 105)
            p2_b.space_after = Pt(10)

    slide2.notes_slide.notes_text_frame.text = speaker_notes

    output_path = "docs/Vendly_Security_Slide.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_security_presentation()
